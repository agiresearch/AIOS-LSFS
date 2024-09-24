from PyPDF2 import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from docx import Document
import openpyxl
import xlwt
import shutil
import os


def update_pdf(filepath, content):
    reader = PdfReader(filepath)
    pages = len(reader.pages)
    width = reader.pages[0].mediabox.width
    height = reader.pages[0].mediabox.height

    c = canvas.Canvas('new_content.pdf', pagesize=(width,height))
    for _ in range(pages):
        c.drawString(100, 500, content)
        c.showPage()

    shutil.move('new_content.pdf', filepath)

def update_docx(filepath,content):
    doc = Document()
    doc.add_paragraph(content)
    doc.save(filepath)

def update_txt(filepath,content):
    with open(filepath, "w") as file:
        file.write(content)

def update_md(filepath,content):
    with open(filepath, "w") as file:
        file.write("# This is the new content")


def update_xlsx(filepath,content):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = content
    workbook.save(filepath)

def update_xls(filepath,content):
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Sheet1")
    sheet.write(0, 0, content)
    workbook.save(filepath)

def update_file(filepath,content):
    _, extension = os.path.splitext(filepath)

    file_dict = {
    '.pdf': update_pdf,
    '.docx': update_docx,
    '.txt': update_docx,
    '.md': update_md,
    '.xls': update_xls,
    '.xlsx': update_xlsx
    }
    if extension in file_dict:
        file_dict[extension](filepath,content)
    else:
        raise Exception(f"\nDon\' support {extension}" )
    


